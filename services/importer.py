"""Import de relevés bancaires CSV / Excel (.xlsx) en deux temps : aperçu, puis validation.

Gère les particularités des exports des banques françaises : séparateur « ; », virgule
décimale, dates JJ/MM/AAAA, encodage Windows-1252, lignes d'en-tête parasites, colonnes
Débit / Crédit séparées.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import sqlite3
import unicodedata
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from finance import to_cents

from .transactions import CategoryGuesser
from .validation import ValidationError, Validator, ensure_exists

MAX_ROWS = 20_000
FIELDS = ("date", "label", "amount", "debit", "credit", "category")
_HASH_RE = re.compile(r"^[0-9a-f]{64}$")


# --------------------------------------------------------------------------- lecture du fichier

def _norm(text) -> str:
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text.lower()).split())


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")  # ne peut pas échouer


def _read_csv(raw: bytes) -> list[list[str]]:
    text = _decode(raw)
    sample = "\n".join(text.splitlines()[:40])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def _read_xlsx(raw: bytes) -> list[list]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("Lecture Excel indisponible : installez openpyxl (pip install openpyxl).") from exc
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # fichier corrompu, protégé par mot de passe...
        raise ValidationError(f"Fichier Excel illisible : {exc}") from exc
    try:
        sheet = wb.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        wb.close()


def read_table(filename: str, raw: bytes) -> list[tuple[int, list]]:
    """Renvoie [(numéro de ligne dans le fichier, cellules)], lignes vides retirées."""
    ext = Path(filename or "").suffix.lower()
    if not raw:
        raise ValidationError("Le fichier est vide.")
    if ext == ".xls":
        raise ValidationError("Format .xls (ancien Excel) non pris en charge : enregistrez le fichier en .xlsx ou .csv.")
    if ext == ".xlsx":
        rows = _read_xlsx(raw)
    elif ext in (".csv", ".txt", ""):
        rows = _read_csv(raw)
    else:
        raise ValidationError("Format non pris en charge : utilisez un fichier .csv ou .xlsx.")
    rows = [(i + 1, r) for i, r in enumerate(rows) if any(c is not None and str(c).strip() for c in r)]
    if not rows:
        raise ValidationError("Aucune ligne exploitable dans le fichier.")
    return rows


# --------------------------------------------------------------------------- colonnes

def _score_header(cells) -> int:
    keys = ("date", "libelle", "montant", "debit", "credit", "description", "operation", "amount", "label")
    return sum(1 for c in cells if any(k in _norm(c) for k in keys))


def find_header_row(rows: list[list]) -> int:
    """Certaines banques ajoutent des lignes d'information avant l'en-tête."""
    for i, row in enumerate(rows[:30]):
        if _score_header(row) >= 2:
            return i
    return 0


def detect_mapping(headers: list[str]) -> dict[str, int | None]:
    names = [_norm(h) for h in headers]
    mapping: dict[str, int | None] = dict.fromkeys(FIELDS)

    def first(pred, exclude=()):
        return next((i for i, n in enumerate(names) if i not in exclude and pred(n)), None)

    mapping["date"] = first(lambda n: "date" in n and "valeur" not in n)
    if mapping["date"] is None:  # attention : l'index 0 est valide, pas de « or » ici
        mapping["date"] = first(lambda n: "date" in n)
    if mapping["date"] is None:
        mapping["date"] = first(lambda n: n in ("jour", "day"))
    used = {mapping["date"]}
    mapping["debit"] = first(lambda n: "debit" in n, used)
    mapping["credit"] = first(lambda n: "credit" in n, used)
    used |= {mapping["debit"], mapping["credit"]}
    mapping["amount"] = first(lambda n: any(k in n for k in ("montant", "amount", "somme")), used)
    used.add(mapping["amount"])
    mapping["category"] = first(lambda n: "categor" in n, used)
    used.add(mapping["category"])
    for keyword in ("libelle", "intitule", "description", "label", "detail", "operation",
                    "beneficiaire", "nature", "memo", "payee", "tiers"):
        mapping["label"] = first(lambda n, k=keyword: k in n, used)
        if mapping["label"] is not None:
            break
    return mapping


def _clean_mapping(mapping: dict | None, width: int) -> dict[str, int | None]:
    clean = {}
    for field in FIELDS:
        value = (mapping or {}).get(field)
        if value in (None, "", -1, "-1"):
            clean[field] = None
            continue
        try:
            idx = int(value)
        except (TypeError, ValueError) as exc:
            raise ValidationError(f"Colonne invalide pour « {field} ».") from exc
        if not 0 <= idx < width:
            raise ValidationError(f"Colonne hors limites pour « {field} ».")
        clean[field] = idx
    errors = {}
    if clean["date"] is None:
        errors["date"] = "Indiquez la colonne contenant la date."
    if clean["label"] is None:
        errors["label"] = "Indiquez la colonne contenant le libellé."
    if clean["amount"] is None and clean["debit"] is None and clean["credit"] is None:
        errors["amount"] = "Indiquez la colonne Montant, ou les colonnes Débit / Crédit."
    if errors:
        raise ValidationError("Correspondance des colonnes incomplète.", errors)
    return clean


# --------------------------------------------------------------------------- valeurs

_DATE_FORMATS = ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%Y/%m/%d", "%d-%m-%y")


def parse_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # numéro de série Excel
        return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
    text = str(value or "").strip().split(" ")[0].split("T")[0]
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"date illisible « {value} »")


def parse_amount(value) -> int | None:
    """Renvoie des centimes, ou None si la cellule est vide."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return to_cents(repr(float(value)))
    text = str(value).strip()
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = re.sub(r"[^\d,.\-+]", "", text)
    if text in ("", "-", "+"):
        raise ValueError(f"montant illisible « {value} »")
    cents = to_cents(text)
    return -abs(cents) if negative else cents


def _cell(row: list, idx: int | None):
    return row[idx] if idx is not None and idx < len(row) else None


def make_hash(account_id: int, day: str, cents: int, label: str, occurrence: int) -> str:
    # Normalisation volontairement minimale et figée : la changer rendrait les anciens imports méconnaissables
    key = f"{account_id}|{day}|{cents}|{' '.join(label.lower().split())}|{occurrence}"
    return hashlib.sha256(key.encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------- aperçu

def preview(conn: sqlite3.Connection, filename: str, raw: bytes, account_id, mapping: dict | None = None) -> dict:
    v = Validator({"account_id": account_id})
    v.integer_id("account_id", required=True)
    account_id = v.result()["account_id"]
    ensure_exists(conn, "accounts", account_id, "Compte")

    table = read_table(filename, raw)
    header_idx = find_header_row([cells for _, cells in table])
    headers = [str(h).strip() if h is not None else f"Colonne {i + 1}" for i, h in enumerate(table[header_idx][1])]
    body = table[header_idx + 1:]
    if len(body) > MAX_ROWS:
        raise ValidationError(f"Fichier trop volumineux : {MAX_ROWS} lignes maximum par import.")
    width = max([len(headers), *(len(r) for _, r in body)]) if body else len(headers)
    headers += [f"Colonne {i + 1}" for i in range(len(headers), width)]

    detected = detect_mapping(headers)
    try:
        cols = _clean_mapping(mapping if mapping else detected, width)
    except ValidationError as exc:
        if not exc.fields:
            raise
        # Colonnes non reconnues : on renvoie les en-têtes pour que l'utilisateur les choisisse
        return {
            "filename": filename, "account_id": account_id, "headers": headers,
            "mapping": {k: v for k, v in (mapping or detected).items() if k in FIELDS},
            "detected_mapping": detected, "mapping_errors": exc.fields, "rows": [],
            "summary": {"total": 0, "ok": 0, "duplicates": 0, "errors": 0},
        }

    categories = {r["name"].lower(): (r["id"], r["name"]) for r in conn.execute("SELECT id, name FROM categories")}
    names = {cid: name for cid, name in categories.values()}
    guesser = CategoryGuesser(conn)
    occurrences: Counter = Counter()
    existing_cache: dict[tuple, int] = {}
    rows, summary = [], Counter()

    for line, raw_row in body:  # numéro de ligne tel qu'affiché dans un tableur
        item = {"line": line, "status": "ok", "error": None}
        try:
            item["date"] = parse_date(_cell(raw_row, cols["date"]))
            label = " ".join(str(_cell(raw_row, cols["label"]) or "").split())
            if not label:
                raise ValueError("libellé vide")
            item["label"] = label[:200]
            if cols["amount"] is not None:
                cents = parse_amount(_cell(raw_row, cols["amount"]))
            else:
                debit = parse_amount(_cell(raw_row, cols["debit"]))
                credit = parse_amount(_cell(raw_row, cols["credit"]))
                cents = (abs(credit) if credit else 0) - (abs(debit) if debit else 0)
                cents = cents if (debit or credit) else None
            if not cents:
                raise ValueError("montant vide ou nul")
            item["amount_cents"] = cents
        except ValueError as exc:
            item.update(status="error", error=str(exc).capitalize(),
                        raw=[("" if c is None else str(c)) for c in raw_row][:width])
            rows.append(item)
            summary["errors"] += 1
            continue

        file_category = str(_cell(raw_row, cols["category"]) or "").strip().lower()
        category_id = categories.get(file_category, (None,))[0] if file_category else None
        category_id = category_id or guesser.guess(item["label"])
        item["category_id"] = category_id
        item["category_name"] = names.get(category_id)

        key = (item["date"], cents, item["label"].lower())
        occurrence = occurrences[key]
        occurrences[key] += 1
        item["import_hash"] = make_hash(account_id, item["date"], cents, item["label"], occurrence)

        if key not in existing_cache:
            existing_cache[key] = conn.execute(
                "SELECT COUNT(*) FROM transactions WHERE account_id = ? AND date = ? AND amount_cents = ? "
                "AND lower(label) = ?",
                (account_id, *key),
            ).fetchone()[0]
        already = conn.execute("SELECT 1 FROM transactions WHERE import_hash = ?", (item["import_hash"],)).fetchone()
        if already or occurrence < existing_cache[key]:
            item["status"] = "duplicate"
            summary["duplicates"] += 1
        else:
            summary["ok"] += 1
        rows.append(item)

    return {
        "filename": filename,
        "account_id": account_id,
        "headers": headers,
        "mapping": cols,
        "detected_mapping": detected,
        "mapping_errors": None,
        "rows": rows,
        "summary": {"total": len(rows), "ok": summary["ok"], "duplicates": summary["duplicates"],
                    "errors": summary["errors"]},
    }


# --------------------------------------------------------------------------- validation

def commit(conn: sqlite3.Connection, data: dict | None) -> dict:
    """Insère les lignes validées dans l'aperçu. Les doublons (même empreinte) sont ignorés."""
    data = data or {}
    v = Validator(data)
    v.integer_id("account_id", required=True)
    account_id = v.result()["account_id"]
    ensure_exists(conn, "accounts", account_id, "Compte")
    rows = data.get("rows")
    if not isinstance(rows, list) or not rows:
        raise ValidationError("Aucune ligne à importer.")
    if len(rows) > MAX_ROWS:
        raise ValidationError(f"{MAX_ROWS} lignes maximum par import.")

    valid_categories = {r[0] for r in conn.execute("SELECT id FROM categories")}
    to_insert, errors = [], {}
    for i, row in enumerate(rows):
        rv = Validator(row if isinstance(row, dict) else None)
        rv.date("date", required=True)
        rv.text("label", required=True, max_len=200)
        rv.integer_id("category_id")
        try:
            clean = rv.result()
            cents = int(row.get("amount_cents"))
            if cents == 0 or abs(cents) >= 10**13:
                raise ValueError
            digest = str(row.get("import_hash") or "")
            if not _HASH_RE.match(digest):
                raise ValueError
        except (ValidationError, TypeError, ValueError, AttributeError):
            errors[str(i)] = "Ligne invalide."
            continue
        category_id = clean.get("category_id") if clean.get("category_id") in valid_categories else None
        to_insert.append((account_id, category_id, clean["date"], clean["label"], cents, digest))
    if errors:
        raise ValidationError(f"{len(errors)} ligne(s) invalide(s) : relancez l'aperçu.", errors)

    before = conn.total_changes
    conn.executemany(
        "INSERT OR IGNORE INTO transactions (account_id, category_id, date, label, amount_cents, source, import_hash) "
        "VALUES (?, ?, ?, ?, ?, 'import', ?)",
        to_insert,
    )
    conn.commit()
    inserted = conn.total_changes - before
    return {"inserted": inserted, "skipped": len(to_insert) - inserted}
